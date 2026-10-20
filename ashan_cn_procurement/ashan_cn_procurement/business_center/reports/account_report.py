#!/usr/bin/env python3
from __future__ import annotations

import base64
import io
import json
import urllib.parse
from .. import local_urllib
from dataclasses import dataclass
from datetime import datetime
from typing import Any
from urllib.error import HTTPError

from .. import legacy_wmill as wmill
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE_URL_VARIABLE_PATH = "f/erpnext/base_url"
API_TOKEN_VARIABLE_PATH = "f/erpnext/api_token"
BACKEND_VERSION = "2.0.0-frappe16"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TIMEOUT_SECONDS = 60
BASE_URL = ""
TOKEN = ""
HEADERS: dict[str, str] = {}

GL_FIELDS = [
    "name", "voucher_type", "voucher_no", "posting_date", "company",
    "account", "against", "debit", "credit", "remarks",
]
COLUMNS = [
    "凭证号", "凭证日期", "类型", "供应商/对方", "付款方式", "项目名(物料名)",
    "规格型号", "单价", "数量", "单位", "发票号码", "分配金额", "借贷方向", "备注",
]


@dataclass(frozen=True)
class Scope:
    company: str
    account: str
    sheet: str


@dataclass
class ScopeResult:
    scope: Scope
    gl_count: int
    row_count: int
    credit_total: float
    debit_total: float
    rows: list[dict[str, Any]]


PE_CACHE: dict[str, dict[str, Any]] = {}
PI_CACHE: dict[str, dict[str, Any]] = {}


def _read_connection_values() -> tuple[str, str, str]:
    """Read the ERPNext connection from the App-managed ERPNext 当前站点适配器.

    The URL is stored at f/erpnext/base_url and the credential at
    f/erpnext/api_token. The token value is never returned to the frontend or
    written to logs.
    """
    variable_errors: list[str] = []

    try:
        base_url = str(wmill.get_variable(BASE_URL_VARIABLE_PATH) or "").strip().rstrip("/")
    except Exception as exc:
        base_url = ""
        variable_errors.append(f"{BASE_URL_VARIABLE_PATH}: {exc}")

    try:
        token = str(wmill.get_variable(API_TOKEN_VARIABLE_PATH) or "").strip()
    except Exception as exc:
        token = ""
        variable_errors.append(f"{API_TOKEN_VARIABLE_PATH}: {exc}")

    missing: list[str] = []
    if not base_url:
        missing.append(BASE_URL_VARIABLE_PATH)
    if not token:
        missing.append(API_TOKEN_VARIABLE_PATH)
    if missing:
        detail = f"；读取错误：{' | '.join(variable_errors)}" if variable_errors else ""
        raise RuntimeError(
            "ERPNext连接配置不可用。请在ERPNext业务中心的“连接设置”页面完成配置："
            + "、".join(missing)
            + detail
        )

    return base_url, token, "frappe_local"

def _set_connection() -> None:
    global BASE_URL, TOKEN, HEADERS
    BASE_URL, TOKEN, _source = _read_connection_values()
    HEADERS = {"Authorization": f"Token {TOKEN}", "Accept": "application/json"}


def fmt_dt(value: datetime) -> str:
    return value.strftime("%Y-%m-%d %H:%M:%S")


def fmt_elapsed(seconds: float) -> str:
    total = int(round(seconds))
    hours, rem = divmod(total, 3600)
    minutes, secs = divmod(rem, 60)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}"


def to_float(value: Any) -> float:
    try:
        return float(value or 0)
    except Exception:
        return 0.0


def load_json(url: str) -> dict[str, Any]:
    request = local_urllib.Request(url, headers=HEADERS)
    with local_urllib.urlopen(request, timeout=TIMEOUT_SECONDS) as response:
        return json.loads(response.read().decode("utf-8"))


def api_get(doctype: str, params: dict[str, str]) -> list[dict[str, Any]]:
    path = f"/api/resource/{urllib.parse.quote(doctype, safe='')}"
    payload = load_json(f"{BASE_URL}{path}?{urllib.parse.urlencode(params)}")
    data = payload.get("data", [])
    return data if isinstance(data, list) else []


def api_doc(doctype: str, name: str) -> dict[str, Any] | None:
    path = f"/api/resource/{urllib.parse.quote(doctype, safe='')}/{urllib.parse.quote(name, safe='')}"
    try:
        payload = load_json(f"{BASE_URL}{path}")
    except HTTPError as exc:
        if exc.code == 404:
            return None
        raise
    data = payload.get("data")
    return data if isinstance(data, dict) else None


def sanitize_sheet_name(name: str) -> str:
    cleaned = "".join("-" if char in "[]:*?/\\" else char for char in str(name or "")).strip().strip("'")
    return (cleaned or "账单明细")[:31]


def unique_sheet_name(base: str, used: set[str]) -> str:
    candidate = sanitize_sheet_name(base)
    original = candidate
    index = 2
    while candidate in used:
        suffix = f"-{index}"
        candidate = f"{original[:31-len(suffix)]}{suffix}"
        index += 1
    used.add(candidate)
    return candidate


def normalize_granularity(value: str) -> str:
    mapping = {"item": "item", "明细": "item", "detail": "item", "voucher": "voucher", "汇总": "voucher", "凭证": "voucher"}
    normalized = str(value or "item").strip().lower()
    if normalized in mapping:
        return mapping[normalized]
    raise ValueError("granularity 只能是 item 或 voucher")


def normalize_direction(value: str) -> str:
    mapping = {"all": "all", "全部": "all", "支出": "支出", "credit": "支出", "收入": "收入", "debit": "收入"}
    normalized = str(value or "all").strip().lower()
    if normalized in mapping:
        return mapping[normalized]
    raise ValueError("direction 只能是 all / 支出 / 收入")


def direction_of(gle: dict[str, Any]) -> str:
    if to_float(gle.get("credit")) > 0:
        return "支出"
    if to_float(gle.get("debit")) > 0:
        return "收入"
    return "-"


def voucher_amount(gle: dict[str, Any]) -> float:
    credit = to_float(gle.get("credit"))
    return credit if credit > 0 else to_float(gle.get("debit"))


def fetch_gl_entries(scope: Scope, start_date: str, end_date: str) -> list[dict[str, Any]]:
    filters = [["company", "=", scope.company], ["account", "=", scope.account], ["posting_date", "between", [start_date, end_date]]]
    rows: list[dict[str, Any]] = []
    start = 0
    page_size = 200
    while True:
        chunk = api_get("GL Entry", {
            "fields": json.dumps(GL_FIELDS, ensure_ascii=False, separators=(",", ":")),
            "filters": json.dumps(filters, ensure_ascii=False, separators=(",", ":")),
            "order_by": "posting_date asc,name asc",
            "limit_start": str(start),
            "limit_page_length": str(page_size),
        })
        rows.extend(chunk)
        if len(chunk) < page_size:
            break
        start += page_size
    return rows


def get_pe(name: str) -> dict[str, Any] | None:
    if name not in PE_CACHE:
        PE_CACHE[name] = api_doc("Payment Entry", name) or {}
    return PE_CACHE[name] or None


def get_pi(name: str) -> dict[str, Any] | None:
    if name not in PI_CACHE:
        PI_CACHE[name] = api_doc("Purchase Invoice", name) or {}
    return PI_CACHE[name] or None


def allocate_by_items(items: list[dict[str, Any]], pi_total: float, allocated_amount: float) -> list[float]:
    if not items:
        return []
    if pi_total <= 0:
        return [0.0 for _ in items]
    raw = [round(to_float(item.get("amount")) * allocated_amount / pi_total, 2) for item in items]
    difference = round(allocated_amount - sum(raw), 2)
    if raw and difference:
        raw[-1] = round(raw[-1] + difference, 2)
    return raw


def make_row(voucher_no: str, posting_date: str, voucher_type: str, against: str, payment_mode: str,
             item_name: str, spec: str, rate: float, qty: float, uom: str, bill_no: str,
             allocated: float, direction: str, remark: str) -> dict[str, Any]:
    return {
        "凭证号": voucher_no, "凭证日期": posting_date, "类型": voucher_type,
        "供应商/对方": against, "付款方式": payment_mode, "项目名(物料名)": item_name,
        "规格型号": spec, "单价": rate, "数量": qty, "单位": uom,
        "发票号码": bill_no, "分配金额": allocated, "借贷方向": direction, "备注": remark,
    }


def brief_rows(gle: dict[str, Any]) -> list[dict[str, Any]]:
    voucher_type = str(gle.get("voucher_type") or "-")
    voucher_no = str(gle.get("voucher_no") or "")
    amount = voucher_amount(gle)
    payment_mode = "-"
    bill_no = "-"
    item_name = "(凭证汇总)"
    if voucher_type == "Payment Entry":
        pe = get_pe(voucher_no)
        if pe:
            payment_mode = str(pe.get("mode_of_payment") or "-")
            refs = [row for row in (pe.get("references") or []) if row.get("reference_doctype") == "Purchase Invoice"]
            if refs:
                first = refs[0]
                pi_name = str(first.get("reference_name") or "")
                pi = get_pi(pi_name) if pi_name else None
                bill_no = str(first.get("bill_no") or (pi or {}).get("bill_no") or pi_name or "-")
    elif voucher_type == "Purchase Invoice":
        pi = get_pi(voucher_no)
        if pi:
            bill_no = str(pi.get("bill_no") or pi.get("name") or voucher_no)
            item_name = str(pi.get("supplier") or item_name)
    return [make_row(
        voucher_no, str(gle.get("posting_date") or ""), voucher_type, str(gle.get("against") or ""),
        payment_mode, item_name, "-", amount, 1, "-", bill_no, amount,
        direction_of(gle), str(gle.get("remarks") or "")[:100],
    )]


def detailed_rows(gle: dict[str, Any]) -> list[dict[str, Any]]:
    voucher_type = str(gle.get("voucher_type") or "-")
    voucher_no = str(gle.get("voucher_no") or "")
    posting_date = str(gle.get("posting_date") or "")
    against = str(gle.get("against") or "")
    amount = voucher_amount(gle)
    direction = direction_of(gle)
    fallback_remark = str(gle.get("remarks") or "")[:100]
    output: list[dict[str, Any]] = []

    if voucher_type == "Payment Entry":
        pe = get_pe(voucher_no)
        if not pe:
            return [make_row(voucher_no, posting_date, voucher_type, against, "-", "(无法获取付款单明细)", "-", 0, 0, "-", "-", amount, direction, fallback_remark)]
        payment_mode = str(pe.get("mode_of_payment") or "-")
        refs = [row for row in (pe.get("references") or []) if row.get("reference_doctype") == "Purchase Invoice"]
        if not refs:
            return [make_row(voucher_no, posting_date, voucher_type, against, payment_mode, "(无 Purchase Invoice 引用)", "-", 0, 0, "-", "-", amount, direction, fallback_remark)]
        for ref in refs:
            pi_name = str(ref.get("reference_name") or "")
            allocated = to_float(ref.get("allocated_amount"))
            pi = get_pi(pi_name) if pi_name else None
            if not pi:
                output.append(make_row(voucher_no, posting_date, voucher_type, against, payment_mode, f"(PI不存在: {pi_name})", "-", 0, 0, "-", pi_name or "-", allocated, direction, fallback_remark))
                continue
            items = pi.get("items") or []
            bill_no = str(ref.get("bill_no") or pi.get("bill_no") or pi.get("name") or pi_name or "-")
            if not items:
                output.append(make_row(voucher_no, posting_date, voucher_type, against, payment_mode, str(pi.get("supplier") or "-"), "-", to_float(pi.get("grand_total")), 1, "-", bill_no, allocated, direction, str(pi.get("custom_备注") or "")))
                continue
            for item, share in zip(items, allocate_by_items(items, to_float(pi.get("grand_total")), allocated)):
                output.append(make_row(
                    voucher_no, posting_date, voucher_type, against, payment_mode,
                    str(item.get("item_name") or item.get("item_code") or "-"),
                    str(item.get("custom_guige_xinghao") or "-"), to_float(item.get("rate")),
                    to_float(item.get("qty")), str(item.get("uom") or "-"), bill_no,
                    share, direction, str(item.get("custom_备注") or ""),
                ))
        return output

    if voucher_type == "Purchase Invoice":
        pi = get_pi(voucher_no)
        if not pi:
            return [make_row(voucher_no, posting_date, voucher_type, against, "-", "(无法获取发票明细)", "-", 0, 0, "-", "-", amount, direction, fallback_remark)]
        items = pi.get("items") or []
        bill_no = str(pi.get("bill_no") or pi.get("name") or voucher_no)
        if not items:
            return [make_row(voucher_no, posting_date, voucher_type, against, "-", str(pi.get("supplier") or "-"), "-", to_float(pi.get("grand_total")), 1, "-", bill_no, amount, direction, str(pi.get("custom_备注") or fallback_remark))]
        for item, share in zip(items, allocate_by_items(items, to_float(pi.get("grand_total")), amount)):
            output.append(make_row(
                voucher_no, posting_date, voucher_type, against, "-",
                str(item.get("item_name") or item.get("item_code") or "-"),
                str(item.get("custom_guige_xinghao") or "-"), to_float(item.get("rate")),
                to_float(item.get("qty")), str(item.get("uom") or "-"), bill_no,
                share, direction, str(item.get("custom_备注") or ""),
            ))
        return output

    return [make_row(voucher_no, posting_date, voucher_type, against, "-", "(非采购类型)", "-", 0, 0, "-", "-", amount, direction, fallback_remark)]


def build_sheet(ws, scope: Scope, rows: list[dict[str, Any]], start_date: str, end_date: str, granularity: str, direction: str) -> tuple[float, float]:
    title_fill = PatternFill("solid", fgColor="17375E")
    header_fill = PatternFill("solid", fgColor="DCE6F1")
    alt_fill = PatternFill("solid", fgColor="F7FAFC")
    total_fill = PatternFill("solid", fgColor="E2F0D9")
    white_bold = Font(bold=True, color="FFFFFF", size=13)
    header_font = Font(bold=True, color="1F2937")
    thin = Side("thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", shrink_to_fit=True)
    left = Alignment(horizontal="left", vertical="center", shrink_to_fit=True)
    right = Alignment(horizontal="right", vertical="center", shrink_to_fit=True)

    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"
    ws.merge_cells("A1:N1")
    ws["A1"] = f"账户明细：{scope.company}｜{scope.account}"
    ws["A1"].font = white_bold
    ws["A1"].fill = title_fill
    ws["A1"].alignment = center
    ws.merge_cells("A2:N2")
    ws["A2"] = f"{start_date} ~ {end_date}｜粒度：{'物料明细' if granularity == 'item' else '凭证汇总'}｜方向：{'全部' if direction == 'all' else direction}｜{len(rows)} 行"
    ws["A2"].alignment = center

    for col, name in enumerate(COLUMNS, 1):
        cell = ws.cell(3, col, name)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = center

    for row_index, row in enumerate(rows, 4):
        for col, name in enumerate(COLUMNS, 1):
            cell = ws.cell(row_index, col, row.get(name, ""))
            cell.border = border
            if (row_index - 4) % 2:
                cell.fill = alt_fill
            if name in {"凭证号", "凭证日期", "类型", "付款方式", "单位", "发票号码", "借贷方向"}:
                cell.alignment = center
            elif name in {"单价", "数量", "分配金额"}:
                cell.alignment = right
            else:
                cell.alignment = left
            if name == "单价":
                cell.number_format = "#,##0.0000"
            elif name == "数量":
                cell.number_format = "#,##0.000"
            elif name == "分配金额":
                cell.number_format = "#,##0.00"
            elif name == "发票号码":
                cell.number_format = "@"

    credit_total = round(sum(to_float(row.get("分配金额")) for row in rows if row.get("借贷方向") == "支出"), 2)
    debit_total = round(sum(to_float(row.get("分配金额")) for row in rows if row.get("借贷方向") == "收入"), 2)
    total_row = len(rows) + 4
    for col in range(1, len(COLUMNS) + 1):
        cell = ws.cell(total_row, col)
        cell.border = border
        cell.fill = total_fill
    ws.cell(total_row, 1, "合计").font = Font(bold=True)
    ws.cell(total_row, 12, credit_total if direction != "收入" else debit_total).font = Font(bold=True)
    ws.cell(total_row, 12).number_format = "#,##0.00"
    ws.cell(total_row, 12).alignment = right

    widths = [22, 12, 16, 20, 12, 24, 18, 12, 10, 9, 18, 14, 10, 26]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    return credit_total, debit_total


def build_info_sheet(ws, started: datetime, ended: datetime, start_date: str, end_date: str,
                     granularity: str, direction: str, results: list[ScopeResult]) -> None:
    title_fill = PatternFill("solid", fgColor="17375E")
    header_fill = PatternFill("solid", fgColor="DCE6F1")
    thin = Side("thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:G1")
    ws["A1"] = "运行信息"
    ws["A1"].font = Font(bold=True, color="FFFFFF", size=13)
    ws["A1"].fill = title_fill
    ws["A1"].alignment = center
    meta = [
        ("脚本开始时间", fmt_dt(started)), ("脚本结束时间", fmt_dt(ended)),
        ("总耗时", fmt_elapsed((ended-started).total_seconds())), ("日期范围", f"{start_date} ~ {end_date}"),
        ("粒度", "物料明细" if granularity == "item" else "凭证汇总"),
        ("方向", "全部" if direction == "all" else direction), ("账户数", len(results)),
    ]
    for row, (key, value) in enumerate(meta, 3):
        ws.cell(row, 1, key).border = border
        ws.cell(row, 2, value).border = border
    header_row = len(meta) + 4
    headers = ["Sheet", "公司", "账户", "GL条数", "输出行数", "支出合计", "收入合计"]
    for col, name in enumerate(headers, 1):
        cell = ws.cell(header_row, col, name)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.border = border
        cell.alignment = center
    for index, result in enumerate(results, header_row + 1):
        values = [result.scope.sheet, result.scope.company, result.scope.account, result.gl_count, result.row_count, result.credit_total, result.debit_total]
        for col, value in enumerate(values, 1):
            cell = ws.cell(index, col, value)
            cell.border = border
            if col in (6, 7):
                cell.number_format = "#,##0.00"
    for index, width in enumerate([18, 24, 30, 10, 10, 14, 14], 1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _short_company(company: str) -> str:
    if "祺富" in company:
        return "祺富"
    if "吉众" in company:
        return "吉众"
    return company[:8]


def _filename(company: str, start_date: str, end_date: str, account_count: int) -> str:
    suffix = "多账户" if account_count > 1 else "账户"
    return f"{_short_company(company)}_{start_date}_{end_date}_{suffix}明细.xlsx"


def main(company: str, accounts: list[str], start_date: str, end_date: str,
         granularity: str = "item", direction: str = "all") -> dict[str, Any]:
    company = str(company or "").strip()
    accounts = [str(value).strip() for value in (accounts or []) if str(value).strip()]
    start_date = str(start_date or "").strip()
    end_date = str(end_date or "").strip()
    if not company:
        raise ValueError("请选择公司")
    if not accounts:
        raise ValueError("至少选择一个账户")
    if not start_date or not end_date or start_date > end_date:
        raise ValueError("日期范围无效")
    granularity = normalize_granularity(granularity)
    direction = normalize_direction(direction)
    _set_connection()
    PE_CACHE.clear()
    PI_CACHE.clear()

    started = datetime.now()
    wb = Workbook()
    info_ws = wb.active
    info_ws.title = "运行信息"
    used = {"运行信息"}
    results: list[ScopeResult] = []
    preview_rows: list[dict[str, Any]] = []

    for account in accounts:
        sheet = unique_sheet_name(account, used)
        scope = Scope(company=company, account=account, sheet=sheet)
        gl_entries = fetch_gl_entries(scope, start_date, end_date)
        rows: list[dict[str, Any]] = []
        for gle in gl_entries:
            generated = brief_rows(gle) if granularity == "voucher" else detailed_rows(gle)
            if direction != "all":
                generated = [row for row in generated if row.get("借贷方向") == direction]
            rows.extend(generated)
        credit, debit = build_sheet(wb.create_sheet(sheet), scope, rows, start_date, end_date, granularity, direction)
        result = ScopeResult(scope, len(gl_entries), len(rows), credit, debit, rows)
        results.append(result)
        for row in rows[: max(0, 100 - len(preview_rows))]:
            preview_rows.append({"账户": account, **row})

    ended = datetime.now()
    build_info_sheet(info_ws, started, ended, start_date, end_date, granularity, direction, results)
    buffer = io.BytesIO()
    wb.save(buffer)
    content = buffer.getvalue()
    verified = load_workbook(io.BytesIO(content), read_only=True)
    filename = _filename(company, start_date, end_date, len(accounts))
    summary = [{
        "Sheet": result.scope.sheet,
        "账户": result.scope.account,
        "GL条数": result.gl_count,
        "输出行数": result.row_count,
        "支出合计": round(result.credit_total, 2),
        "收入合计": round(result.debit_total, 2),
    } for result in results]
    markdown = (
        f"### 账户明细导出完成\n\n"
        f"- 公司：{company}\n- 日期：{start_date} 至 {end_date}\n"
        f"- 账户：{len(accounts)} 个\n- 输出行：{sum(item.row_count for item in results)} 行\n"
        f"- 支出合计：{sum(item.credit_total for item in results):,.2f}\n"
        f"- 收入合计：{sum(item.debit_total for item in results):,.2f}\n"
        f"- 工作表：{'、'.join(verified.sheetnames)}"
    )
    encoded = base64.b64encode(content).decode("ascii")
    return {
        "render_all": [
            {"markdown": markdown},
            {"table-row-object": summary},
            {"file": {"filename": filename, "content": encoded, "mime_type": XLSX_MIME}},
        ],
        "summary": summary,
        "rows": preview_rows,
        "filename": filename,
        "sheet_names": verified.sheetnames,
        "backend_version": BACKEND_VERSION,
    }
