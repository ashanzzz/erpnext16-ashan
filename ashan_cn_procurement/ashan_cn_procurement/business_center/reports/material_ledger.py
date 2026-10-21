#!/usr/bin/env python3
"""Self-contained ERPNext monthly material ledger generator.

This version no longer depends on an external frozen template file at runtime.
Instead it reconstructs the approved workbook layout directly in openpyxl so the
result stays compatible with the approved Gizhong v0.31 template shape.
"""
from __future__ import annotations

import base64
import io
import json
import re
import urllib.error
import urllib.parse
from .. import local_urllib
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from typing import Any

from .. import legacy_wmill as wmill

from openpyxl import Workbook, load_workbook  # type: ignore
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.hyperlink import Hyperlink

# --- Config -----------------------------------------------------------------
BASE_URL_VARIABLE_PATH = 'f/erpnext/base_url'
API_TOKEN_VARIABLE_PATH = 'f/erpnext/api_token'
SUMMARY_SHEET_NAME = '汇总'
DEFAULT_COMPANY = '天津吉众科技有限公司'
DEFAULT_PARENT_GROUP = '原材料'
SUMMARY_COLS = 14
DETAIL_COLS = 10
DATE_FMT = 'yyyy-mm-dd'
NUM_FMT_DEC = '#,##0.00'
API_TIMEOUT_SECONDS = 60
BACKEND_VERSION = "2.0.0-frappe16"

SUMMARY_HEADERS = [
    '物料编号', '物料名称', '规格型号', '父物料组', '子物料组',
    '上期结转数量', '上期结转金额',
    '本期入库数量', '本期入库金额',
    '本期出库数量', '本期出库金额',
    '本期结余数量', '本期结余金额',
    '备注/状态',
]
DETAIL_HEADERS = ['日期', '凭证类型', '凭证号', '入库数量', '入库金额', '出库数量', '出库金额', '结余数量', '结余金额', '备注']

# --- Style palette ----------------------------------------------------------
FONT_REG = Font(name='Microsoft YaHei', size=10, color='00000000')
FONT_BOLD = Font(name='Microsoft YaHei', size=10, bold=True, color='00000000')
FILL_TITLE = PatternFill(fill_type='solid', fgColor='00D9EAF7')
FILL_LIGHT = PatternFill(fill_type='solid', fgColor='00F2F2F2')
FILL_HDR_BLUE = PatternFill(fill_type='solid', fgColor='00EAF3FB')
FILL_GREEN = PatternFill(fill_type='solid', fgColor='00E2F0D9')
FILL_ORANGE = PatternFill(fill_type='solid', fgColor='00FCE4D6')
FILL_PURPLE = PatternFill(fill_type='solid', fgColor='00E4DFEC')
FILL_YELLOW = PatternFill(fill_type='solid', fgColor='00FFF2CC')
FILL_WHITE = PatternFill(fill_type='solid', fgColor='00FFFFFF')
THIN = Side('thin')
MEDIUM = Side('medium')
BORDER_THIN = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BORDER_MEDIUM = Border(left=MEDIUM, right=MEDIUM, top=MEDIUM, bottom=MEDIUM)
ALIGN_CENTER = Alignment(horizontal='center', vertical='center')
ALIGN_CENTER_SHRINK = Alignment(horizontal='center', vertical='center', shrink_to_fit=True)
ALIGN_LEFT = Alignment(horizontal='left', vertical='center')

STYLE_SPECS: dict[str, tuple[Any, Any, Any, Any, str]] = {
    'title': (FONT_BOLD, FILL_TITLE, BORDER_MEDIUM, ALIGN_LEFT, 'General'),
    'meta_label': (FONT_BOLD, FILL_LIGHT, BORDER_THIN, ALIGN_CENTER, 'General'),
    'meta_value': (FONT_REG, FILL_LIGHT, BORDER_THIN, ALIGN_CENTER, 'General'),
    'summary_head_blue': (FONT_BOLD, FILL_HDR_BLUE, BORDER_THIN, ALIGN_CENTER, 'General'),
    'summary_head_green': (FONT_BOLD, FILL_GREEN, BORDER_THIN, ALIGN_CENTER, 'General'),
    'summary_head_orange': (FONT_BOLD, FILL_ORANGE, BORDER_THIN, ALIGN_CENTER, 'General'),
    'summary_head_purple': (FONT_BOLD, FILL_PURPLE, BORDER_THIN, ALIGN_CENTER, 'General'),
    'summary_head_green_dec': (FONT_BOLD, FILL_GREEN, BORDER_THIN, ALIGN_CENTER, NUM_FMT_DEC),
    'summary_head_orange_dec': (FONT_BOLD, FILL_ORANGE, BORDER_THIN, ALIGN_CENTER, NUM_FMT_DEC),
    'summary_body': (FONT_REG, FILL_WHITE, BORDER_THIN, ALIGN_CENTER, 'General'),
    'summary_num_green_dec': (FONT_REG, FILL_GREEN, BORDER_THIN, ALIGN_CENTER, NUM_FMT_DEC),
    'summary_num_orange_dec': (FONT_REG, FILL_ORANGE, BORDER_THIN, ALIGN_CENTER, NUM_FMT_DEC),
    'summary_num_green_plain': (FONT_REG, FILL_GREEN, BORDER_THIN, ALIGN_CENTER, NUM_FMT_DEC),
    'summary_num_orange_plain': (FONT_REG, FILL_ORANGE, BORDER_THIN, ALIGN_CENTER, NUM_FMT_DEC),
    'summary_status': (FONT_REG, FILL_PURPLE, BORDER_THIN, ALIGN_CENTER, 'General'),
    'detail_link': (FONT_BOLD, FILL_YELLOW, BORDER_MEDIUM, ALIGN_CENTER, 'General'),
    'detail_note': (FONT_REG, FILL_LIGHT, BORDER_THIN, ALIGN_LEFT, 'General'),
    'detail_header': (FONT_BOLD, FILL_HDR_BLUE, BORDER_MEDIUM, ALIGN_CENTER, 'General'),
    'detail_open_label': (FONT_BOLD, FILL_YELLOW, BORDER_THIN, ALIGN_CENTER, 'General'),
    'detail_open_qty': (FONT_BOLD, FILL_GREEN, BORDER_THIN, ALIGN_CENTER, 'General'),
    'detail_open_amt': (FONT_BOLD, FILL_ORANGE, BORDER_THIN, ALIGN_CENTER, 'General'),
    'detail_text': (FONT_REG, FILL_WHITE, BORDER_THIN, ALIGN_CENTER, 'General'),
    'detail_num_green_dec': (FONT_REG, FILL_GREEN, BORDER_THIN, ALIGN_CENTER, NUM_FMT_DEC),
    'detail_num_orange_dec': (FONT_REG, FILL_ORANGE, BORDER_THIN, ALIGN_CENTER, NUM_FMT_DEC),
    'detail_num_green_plain': (FONT_REG, FILL_GREEN, BORDER_THIN, ALIGN_CENTER, 'General'),
    'detail_num_orange_plain': (FONT_REG, FILL_ORANGE, BORDER_THIN, ALIGN_CENTER, 'General'),
    'detail_zero': (FONT_REG, FILL_WHITE, BORDER_THIN, ALIGN_CENTER, NUM_FMT_DEC),
    'detail_total_highlight': (FONT_BOLD, FILL_YELLOW, BORDER_MEDIUM, ALIGN_CENTER, 'General'),
    'detail_total_green': (FONT_BOLD, FILL_GREEN, BORDER_MEDIUM, ALIGN_CENTER, 'General'),
    'detail_total_orange': (FONT_BOLD, FILL_ORANGE, BORDER_MEDIUM, ALIGN_CENTER, 'General'),
    'detail_total_purple': (FONT_REG, FILL_PURPLE, BORDER_MEDIUM, ALIGN_CENTER, 'General'),
}


def set_style(cell, key: str) -> None:
    font, fill, border, alignment, numfmt = STYLE_SPECS[key]
    cell.font = font
    cell.fill = fill
    cell.border = border
    cell.alignment = alignment
    cell.number_format = numfmt


# --- Helpers ----------------------------------------------------------------

def api_get(base: str, token: str, doctype: str, params: dict[str, str]) -> dict:
    path = f"/api/resource/{urllib.parse.quote(doctype)}"
    qs = urllib.parse.urlencode(params, doseq=False)
    url = f"{base.rstrip('/')}{path}?{qs}"
    req = local_urllib.Request(url, headers={'Authorization': f'token {token}'})
    try:
        with local_urllib.urlopen(req, timeout=API_TIMEOUT_SECONDS) as r:
            return json.loads(r.read().decode('utf-8'))
    except urllib.error.HTTPError as exc:
        body = exc.read().decode('utf-8', errors='replace')[:1000]
        if exc.code == 401:
            raise RuntimeError('ERPNext 返回 401，请检查 f/erpnext/api_token') from exc
        if exc.code == 403:
            raise RuntimeError(f'ERPNext 返回 403，当前 API 用户无权读取 {doctype}') from exc
        if exc.code == 404:
            raise RuntimeError(f'ERPNext 返回 404，请检查 f/erpnext/base_url。请求类型：{doctype}') from exc
        raise RuntimeError(f'ERPNext API 请求失败，HTTP {exc.code}，类型：{doctype}，响应：{body}') from exc
    except TimeoutError as exc:
        raise RuntimeError(f'ERPNext API timeout after {API_TIMEOUT_SECONDS}s while requesting {doctype}') from exc
    except urllib.error.URLError as exc:
        reason = getattr(exc, 'reason', None)
        if isinstance(reason, TimeoutError):
            raise RuntimeError(f'ERPNext API timeout after {API_TIMEOUT_SECONDS}s while requesting {doctype}') from exc
        raise RuntimeError(f'无法连接 ERPNext，类型：{doctype}，原因：{reason or exc}') from exc


def api_get_all(base: str, token: str, doctype: str, params: dict[str, str], page_size: int = 5000) -> list[dict]:
    rows: list[dict] = []
    start = 0
    while True:
        page_params = dict(params)
        page_params['limit_page_length'] = str(page_size)
        page_params['limit_start'] = str(start)
        resp = api_get(base, token, doctype, page_params)
        batch = resp.get('data') or []
        if not isinstance(batch, list):
            raise RuntimeError(f'{doctype} API response data is not a list')
        rows.extend(batch)
        if len(batch) < page_size:
            break
        start += page_size
    return rows


def normalize_period(raw: str) -> str:
    value = str(raw or '').strip()
    compact = re.fullmatch(r'(\d{8})-(\d{8})', value)
    if compact:
        a, b = compact.groups()
        return f'{a[:4]}-{a[4:6]}-{a[6:]}:{b[:4]}-{b[4:6]}-{b[6:]}'
    return value


def parse_period(raw: str) -> tuple[date, date]:
    value = normalize_period(raw)
    try:
        if ':' in value:
            start_str, end_str = value.split(':', 1)
            start = datetime.strptime(start_str.strip(), '%Y-%m-%d').date()
            end = datetime.strptime(end_str.strip(), '%Y-%m-%d').date()
        else:
            matched = re.fullmatch(r'(\d{4})-(\d{1,2})', value)
            if not matched:
                raise ValueError
            y, m = map(int, matched.groups())
            if not 1 <= m <= 12:
                raise ValueError
            start = date(y, m, 1)
            next_month = date(y + 1, 1, 1) if m == 12 else date(y, m + 1, 1)
            end = next_month - timedelta(days=1)
    except ValueError as exc:
        raise ValueError(
            'period 格式错误。请填写 YYYY-MM、YYYY-MM-DD:YYYY-MM-DD，'
            '或紧凑范围 YYYYMMDD-YYYYMMDD。'
        ) from exc
    if start > end:
        raise ValueError('period 的开始日期不能晚于结束日期')
    return start, end


def to_float(v) -> float:
    if v in (None, ''):
        return 0.0
    return float(v)


def chinese_prefix(text: str, n: int = 4) -> str:
    chars = [ch for ch in str(text) if '\u4e00' <= ch <= '\u9fff']
    if len(chars) >= n:
        return ''.join(chars[:n])
    compact = ''.join(ch for ch in str(text) if ch.strip())
    return compact[:n] or '物料'


def infer_company_short(company_name: str) -> str:
    text = ''.join(ch for ch in str(company_name) if '\u4e00' <= ch <= '\u9fff')
    prefixes = (
        '北京市', '北京', '天津市', '天津', '上海市', '上海', '重庆市', '重庆',
        '深圳市', '深圳', '广州市', '广州', '杭州市', '杭州', '南京市', '南京',
        '成都市', '成都', '苏州市', '苏州', '无锡市', '无锡',
    )
    for p in prefixes:
        if text.startswith(p) and len(text) > len(p):
            text = text[len(p):]
            break
    keywords = ('科技', '实业', '商贸', '工贸', '电子', '化工', '材料', '设备', '贸易', '物流', '服务', '机械', '制造', '信息', '网络', '集团', '发展', '有限责任公司', '有限公司', '公司')
    for kw in keywords:
        idx = text.find(kw)
        if idx > 0:
            return text[:idx][:4]
    return text[:4]


def normalize_sheet_name(item_name: str, used: set[str]) -> str:
    base = str(item_name or '').strip() or '物料'
    base = base[:31]
    if base not in used:
        used.add(base)
        return base
    idx = 2
    while True:
        suffix = str(idx)
        cand = f"{base[:31 - len(suffix)]}{suffix}"[:31]
        if cand not in used:
            used.add(cand)
            return cand
        idx += 1


def sheet_location(sheet_name: str, cell_ref: str = 'A1') -> str:
    escaped_sheet_name = sheet_name.replace("'", "''")
    return "'{}'!{}".format(escaped_sheet_name, cell_ref)


def set_summary_identity_fit(ws, row: int) -> None:
    """Apply shrink-to-fit to material code, name and specification cells."""
    for col in (1, 2, 3):
        ws.cell(row, col).alignment = ALIGN_CENTER_SHRINK


def set_summary_detail_links(ws, row: int, item: 'Item') -> None:
    """Link material identity cells to the matching detail worksheet."""
    location = sheet_location(item.sheet_name)
    for col in (1, 2, 3):
        cell = ws.cell(row, col)
        if cell.value not in (None, ''):
            cell.hyperlink = Hyperlink(
                ref=cell.coordinate,
                location=location,
                display=str(cell.value),
            )


# --- Data classes -----------------------------------------------------------
@dataclass(slots=True)
class Txn:
    date: date
    time: str
    voucher_type: str
    voucher_no: str
    warehouse: str
    in_qty: float
    out_qty: float
    qty_after: float
    stock_value_diff: float
    stock_value: float
    valuation_rate: float


@dataclass(slots=True)
class Item:
    item_code: str
    item_name: str
    item_group: str
    parent_group: str
    spec: str
    uom: str
    rows: list[Txn] = field(default_factory=list)
    current_rows: list[Txn] = field(default_factory=list)
    opening_qty: float = 0.0
    opening_amt: float = 0.0
    inbound_qty: float = 0.0
    inbound_amt: float = 0.0
    outbound_qty: float = 0.0
    outbound_amt: float = 0.0
    ending_qty: float = 0.0
    ending_amt: float = 0.0
    sheet_name: str = ''
    has_current: bool = False


# --- ERPNext data fetching --------------------------------------------------
def fetch_erp_data(period: str, company: str, erpnext: dict[str, Any], start: date, end: date) -> dict[str, Any]:
    base_url = str(erpnext.get('base_url') or '').strip().rstrip('/')
    token = str(erpnext.get('api_token') or '').strip()
    if not base_url or not token:
        raise ValueError('ERPNext Variables 缺少 base_url 或 api_token')
    export_start = date(2000, 1, 1)

    group_rows = api_get_all(base_url, token, 'Item Group', {
        'fields': json.dumps(['name', 'parent_item_group', 'is_group'], separators=(',', ':')),
        'order_by': 'name asc',
    })
    group_parent_map: dict[str, str] = {}
    for g in group_rows:
        name = str(g.get('name') or '').strip()
        parent = str(g.get('parent_item_group') or '').strip()
        if name:
            group_parent_map[name] = parent

    item_rows = api_get_all(base_url, token, 'Item', {
        'fields': json.dumps(['name', 'item_name', 'item_group', 'custom_guige_xinghao', 'stock_uom', 'is_stock_item', 'disabled'], separators=(',', ':')),
        'filters': json.dumps([['disabled', '=', 0], ['is_stock_item', '=', 1]], separators=(',', ':')),
        'order_by': 'name asc',
    })

    sle_rows = api_get_all(base_url, token, 'Stock Ledger Entry', {
        'fields': json.dumps([
            'name', 'posting_date', 'posting_time', 'item_code', 'warehouse', 'actual_qty',
            'qty_after_transaction', 'stock_value_difference', 'stock_value', 'voucher_type',
            'voucher_no', 'company', 'valuation_rate', 'stock_uom'
        ], separators=(',', ':')),
        'filters': json.dumps([
            ['company', '=', company],
            ['posting_date', '>=', export_start.isoformat()],
            ['posting_date', '<=', end.isoformat()],
        ], separators=(',', ':')),
        'order_by': 'posting_date asc, posting_time asc, name asc',
    })

    return {
        'period': period,
        'company': company,
        'group_parent_map': group_parent_map,
        'items': item_rows,
        'sle_rows': sle_rows,
        'start': start.isoformat(),
        'end': end.isoformat(),
        'export_start': export_start.isoformat(),
    }


def build_items(payload: dict[str, Any], start: date, end: date, parent_group_filter: str = '') -> list[Item]:
    group_parent_map: dict[str, str] = payload.get('group_parent_map', {})
    item_rows = payload.get('items', [])
    sle_rows = payload.get('sle_rows', [])

    def get_descendants(root: str) -> set[str]:
        result = {root}
        for name, parent in group_parent_map.items():
            if parent == root:
                result.update(get_descendants(name))
        return result

    allowed_groups = get_descendants(parent_group_filter) if parent_group_filter else set()

    items: dict[str, Item] = {}
    for row in item_rows:
        code = str(row.get('name') or '').strip()
        if not code:
            continue
        ig = str(row.get('item_group') or '').strip()
        parent_ig = group_parent_map.get(ig) or ig
        if parent_group_filter and ig not in allowed_groups:
            continue
        items[code] = Item(
            item_code=code,
            item_name=str(row.get('item_name') or '').strip(),
            item_group=ig,
            parent_group=parent_ig,
            spec=str(row.get('custom_guige_xinghao') or '').strip(),
            uom=str(row.get('stock_uom') or '').strip(),
        )

    for r in sle_rows:
        code = str(r.get('item_code') or '').strip()
        if not code or code not in items:
            continue
        item = items[code]
        dt_str = str(r.get('posting_date') or '')
        if not dt_str:
            continue
        try:
            dt = datetime.strptime(dt_str[:10], '%Y-%m-%d').date()
        except ValueError:
            continue
        tx = Txn(
            date=dt,
            time=str(r.get('posting_time') or ''),
            voucher_type=str(r.get('voucher_type') or ''),
            voucher_no=str(r.get('voucher_no') or ''),
            warehouse=str(r.get('warehouse') or ''),
            in_qty=max(to_float(r.get('actual_qty')), 0.0),
            out_qty=max(-to_float(r.get('actual_qty')), 0.0),
            qty_after=to_float(r.get('qty_after_transaction')),
            stock_value_diff=to_float(r.get('stock_value_difference')),
            stock_value=to_float(r.get('stock_value')),
            valuation_rate=to_float(r.get('valuation_rate')),
        )
        item.rows.append(tx)
        if start <= dt <= end:
            item.current_rows.append(tx)
            item.has_current = True

    for item in items.values():
        item.rows.sort(key=lambda x: (x.date, x.time, x.voucher_no))
        item.current_rows.sort(key=lambda x: (x.date, x.time, x.voucher_no))
        prev_rows = [x for x in item.rows if x.date < start]
        if prev_rows:
            last_prev = prev_rows[-1]
            item.opening_qty = last_prev.qty_after
            item.opening_amt = last_prev.stock_value
        elif item.current_rows:
            first = item.current_rows[0]
            item.opening_qty = first.qty_after - (first.in_qty - first.out_qty)
            item.opening_amt = first.stock_value - first.stock_value_diff

        for tx in item.current_rows:
            item.inbound_qty += tx.in_qty
            item.outbound_qty += tx.out_qty
            if tx.stock_value_diff >= 0:
                item.inbound_amt += tx.stock_value_diff
            else:
                item.outbound_amt += abs(tx.stock_value_diff)

        if item.current_rows:
            last = item.current_rows[-1]
            item.ending_qty = last.qty_after
            item.ending_amt = last.stock_value
        else:
            item.ending_qty = item.opening_qty
            item.ending_amt = item.opening_amt

    return list(items.values())


# --- Workbook construction --------------------------------------------------
def build_summary_template_sheet(ws) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'A5'
    ws.sheet_format.defaultRowHeight = 24.0
    ws.sheet_format.customHeight = True
    ws.page_margins.left = 0.75
    ws.page_margins.right = 0.75
    ws.page_margins.top = 1.0
    ws.page_margins.bottom = 1.0
    ws.page_margins.header = 0.5
    ws.page_margins.footer = 0.5
    widths = {
        'A': 14.0, 'B': 13.375, 'C': 22.0, 'D': 15.0,
        'E': 15.0, 'F': 12.0, 'G': 14.0, 'H': 12.0,
        'I': 14.0, 'J': 12.0, 'K': 14.0, 'L': 12.0,
        'M': 14.0, 'N': 18.0,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for r in [1, 2, 3, 4, 5, 6, 7, 9, 10, 11, 12, 13]:
        ws.row_dimensions[r].height = 24.0


def build_detail_template_sheet(ws) -> None:
    ws.sheet_view.showGridLines = False
    ws.sheet_format.defaultRowHeight = 15.0
    ws.page_margins.left = 0.75
    ws.page_margins.right = 0.75
    ws.page_margins.top = 1.0
    ws.page_margins.bottom = 1.0
    ws.page_margins.header = 0.5
    ws.page_margins.footer = 0.5
    widths = {
        'A': 14.0, 'B': 12.0, 'C': 16.0, 'D': 14.0,
        'E': 13.0, 'F': 13.0, 'G': 13.0, 'H': 13.0,
        'I': 13.0, 'J': 18.0, 'K': 13.0, 'L': 13.0,
        'M': 13.0, 'N': 13.0, 'O': 13.0,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for r in [4, 5, 6, 7]:
        ws.row_dimensions[r].height = 13.5


DETAIL_TEMPLATE_SHEET = '__detail_template__'


def build_template_workbook() -> Workbook:
    """Create the workbook skeleton in code, without reading an external template file."""
    wb = Workbook()
    wb.remove(wb.active)
    summary_ws = wb.create_sheet(SUMMARY_SHEET_NAME, 0)
    build_summary_template_sheet(summary_ws)
    detail_ws = wb.create_sheet(DETAIL_TEMPLATE_SHEET)
    build_detail_template_sheet(detail_ws)
    detail_ws.sheet_state = 'hidden'
    return wb


def summary_header_style(col: int, dec_variant: bool) -> str:
    if col <= 5:
        return 'summary_head_blue'
    if col in (6, 8, 10, 12):
        return 'summary_head_green_dec' if dec_variant else 'summary_head_green'
    if col in (7, 9, 11, 13):
        return 'summary_head_orange_dec' if dec_variant else 'summary_head_orange'
    return 'summary_head_purple'


def summary_body_style(col: int, dec_variant: bool) -> str:
    if col <= 5:
        return 'summary_body'
    if col in (6, 8, 10, 12):
        return 'summary_num_green_dec' if dec_variant else 'summary_num_green_plain'
    if col in (7, 9, 11, 13):
        return 'summary_num_orange_dec' if dec_variant else 'summary_num_orange_plain'
    return 'summary_status'


def format_status(item: Item) -> str:
    if not item.has_current:
        return '本期无活动'
    if abs(item.ending_qty) < 1e-9:
        return '本期结清'
    return '本期有活动'


def build_summary_sheet(ws, company: str, parent_group: str, start: date, end: date, active: list[Item], inactive: list[Item]) -> None:
    build_summary_template_sheet(ws)

    ws['A1'] = '基本信息'
    set_style(ws['A1'], 'title')
    ws.merge_cells('A1:N1')

    ws['A2'] = '公司名称'
    set_style(ws['A2'], 'meta_label')
    ws['B2'] = company
    set_style(ws['B2'], 'meta_value')
    ws.merge_cells('B2:D2')

    ws['E2'] = '筛选父物料组'
    set_style(ws['E2'], 'meta_label')
    ws['F2'] = parent_group or '全部'
    set_style(ws['F2'], 'meta_value')
    ws.merge_cells('F2:I2')

    ws['J2'] = '日期范围'
    set_style(ws['J2'], 'meta_label')
    ws['K2'] = f'{start:%Y年%m月%d日} 至 {end:%Y年%m月%d日}'
    set_style(ws['K2'], 'meta_value')
    ws.merge_cells('K2:N2')

    ws['A3'] = '物料编号、物料名称和规格型号可点击跳转到对应明细页；汇总区分父物料组/子物料组；明细页右上角可返回汇总。'
    set_style(ws['A3'], 'meta_value')
    ws.merge_cells('A3:N3')

    ws['A4'] = '本期有动态'
    set_style(ws['A4'], 'title')
    ws.merge_cells('A4:N4')

    for col, header in enumerate(SUMMARY_HEADERS, 1):
        cell = ws.cell(5, col)
        cell.value = header
        set_style(cell, summary_header_style(col, dec_variant=False))
    set_summary_identity_fit(ws, 5)

    active_start_row = 6

    # Only create rows for actual active materials.
    # The previous implementation reserved at least five body rows with
    # max(5, len(active)), which left one or more visually empty rows before
    # the next section whenever there were fewer than five active materials.
    row = active_start_row
    for idx, item in enumerate(active):
        dec_variant = idx < 2
        for col in range(1, SUMMARY_COLS + 1):
            cell = ws.cell(row, col)
            set_style(cell, summary_body_style(col, dec_variant))
        ws.cell(row, 1).value = item.item_code
        ws.cell(row, 2).value = item.item_name
        ws.cell(row, 3).value = item.spec or None
        ws.cell(row, 4).value = item.parent_group or None
        ws.cell(row, 5).value = item.item_group or None
        ws.cell(row, 6).value = item.opening_qty
        ws.cell(row, 7).value = item.opening_amt
        ws.cell(row, 8).value = item.inbound_qty
        ws.cell(row, 9).value = item.inbound_amt
        ws.cell(row, 10).value = item.outbound_qty
        ws.cell(row, 11).value = item.outbound_amt
        ws.cell(row, 12).value = item.ending_qty
        ws.cell(row, 13).value = item.ending_amt
        ws.cell(row, 14).value = format_status(item)
        set_summary_identity_fit(ws, row)
        set_summary_detail_links(ws, row, item)
        row += 1

    # Place the next section immediately after the final active material.
    # If there are no active materials, it starts directly below row 5.
    inactive_title_row = row
    inactive_header_row = inactive_title_row + 1
    inactive_start_row = inactive_header_row + 1

    ws.cell(inactive_title_row, 1).value = '有结余 / 无动态 物料汇总'
    set_style(ws.cell(inactive_title_row, 1), 'title')
    ws.merge_cells(start_row=inactive_title_row, start_column=1, end_row=inactive_title_row, end_column=14)

    for col, header in enumerate(SUMMARY_HEADERS, 1):
        cell = ws.cell(inactive_header_row, col)
        cell.value = header
        set_style(cell, summary_header_style(col, dec_variant=True))
    set_summary_identity_fit(ws, inactive_header_row)

    row = inactive_start_row
    for idx, item in enumerate(inactive):
        dec_variant = idx == 0
        for col in range(1, SUMMARY_COLS + 1):
            cell = ws.cell(row, col)
            set_style(cell, summary_body_style(col, dec_variant))
        ws.cell(row, 1).value = item.item_code
        ws.cell(row, 2).value = item.item_name
        ws.cell(row, 3).value = item.spec or None
        ws.cell(row, 4).value = item.parent_group or None
        ws.cell(row, 5).value = item.item_group or None
        ws.cell(row, 6).value = item.opening_qty
        ws.cell(row, 7).value = item.opening_amt
        ws.cell(row, 8).value = item.inbound_qty
        ws.cell(row, 9).value = item.inbound_amt
        ws.cell(row, 10).value = item.outbound_qty
        ws.cell(row, 11).value = item.outbound_amt
        ws.cell(row, 12).value = item.ending_qty
        ws.cell(row, 13).value = item.ending_amt
        ws.cell(row, 14).value = '本期无活动'
        set_summary_identity_fit(ws, row)
        row += 1


def populate_detail_tx_row(ws, row: int, tx: Txn, idx: int) -> None:
    first = idx == 0
    # text columns
    for col in (1, 2, 3, 10):
        set_style(ws.cell(row, col), 'detail_text' if col != 10 else 'detail_note')
    # numeric columns
    if first:
        qty_style_name = 'detail_num_green_dec'
        amt_style_name = 'detail_num_orange_dec'
    else:
        qty_style_name = 'detail_num_green_plain'
        amt_style_name = 'detail_num_orange_plain'
    set_style(ws.cell(row, 4), qty_style_name)
    set_style(ws.cell(row, 5), amt_style_name)
    if abs(tx.out_qty) < 1e-9 and abs(tx.stock_value_diff) < 1e-9:
        set_style(ws.cell(row, 6), 'detail_zero')
        set_style(ws.cell(row, 7), 'detail_zero')
    else:
        set_style(ws.cell(row, 6), 'detail_text')
        set_style(ws.cell(row, 7), 'detail_text')
    set_style(ws.cell(row, 8), 'detail_num_green_dec')
    set_style(ws.cell(row, 9), 'detail_num_orange_dec')

    ws.cell(row, 1).value = tx.date
    ws.cell(row, 1).number_format = DATE_FMT
    ws.cell(row, 2).value = tx.voucher_type
    ws.cell(row, 3).value = tx.voucher_no
    ws.cell(row, 4).value = tx.in_qty
    ws.cell(row, 5).value = tx.stock_value_diff if tx.stock_value_diff > 0 else 0
    ws.cell(row, 6).value = tx.out_qty
    ws.cell(row, 7).value = abs(tx.stock_value_diff) if tx.stock_value_diff < 0 else 0
    ws.cell(row, 8).value = tx.qty_after
    ws.cell(row, 9).value = tx.stock_value
    ws.cell(row, 10).value = tx.warehouse


def build_detail_sheet(wb, item: Item, start: date, end: date) -> None:
    ws = wb.create_sheet(item.sheet_name)
    build_detail_template_sheet(ws)

    ws['A1'] = f'{item.item_code} / {item.item_name} 明细'
    set_style(ws['A1'], 'title')
    ws.merge_cells('A1:H1')

    # Use two compatible internal-link mechanisms for the return button.
    # The formula is broadly compatible with Excel/WPS, while the explicit
    # hyperlink location acts as a fallback for viewers that do not recalculate.
    return_location = sheet_location(SUMMARY_SHEET_NAME)
    return_formula_location = '#' + return_location
    ws['I1'] = '=HYPERLINK("{}","返回汇总")'.format(return_formula_location)
    set_style(ws['I1'], 'detail_link')
    ws.merge_cells('I1:J1')
    ws['I1'].hyperlink = Hyperlink(
        ref='I1',
        location=return_location,
        tooltip='返回主工作表：{}'.format(SUMMARY_SHEET_NAME),
        display='返回汇总',
    )

    ws['A2'] = '右上角可返回汇总；汇总中的物料编号、物料名称和规格型号可跳转到本页。'
    set_style(ws['A2'], 'detail_note')
    ws.merge_cells('A2:J2')

    ws['A4'] = '物料编号'
    set_style(ws['A4'], 'meta_label')
    ws['C4'] = item.item_code
    set_style(ws['C4'], 'detail_text')
    ws['F4'] = '父物料组'
    set_style(ws['F4'], 'meta_label')
    ws['H4'] = item.parent_group or None
    set_style(ws['H4'], 'detail_text')
    ws.merge_cells('A4:B4')
    ws.merge_cells('C4:E4')
    ws.merge_cells('F4:G4')
    ws.merge_cells('H4:J4')

    ws['A5'] = '物料名称'
    set_style(ws['A5'], 'meta_label')
    ws['C5'] = item.item_name
    set_style(ws['C5'], 'detail_text')
    ws['F5'] = '子物料组'
    set_style(ws['F5'], 'meta_label')
    ws['H5'] = item.item_group or None
    set_style(ws['H5'], 'detail_text')
    ws.merge_cells('A5:B5')
    ws.merge_cells('C5:E5')
    ws.merge_cells('F5:G5')
    ws.merge_cells('H5:J5')

    ws['A6'] = '单位'
    set_style(ws['A6'], 'meta_label')
    ws['C6'] = item.uom or None
    set_style(ws['C6'], 'detail_text')
    ws['F6'] = '规格型号'
    set_style(ws['F6'], 'meta_label')
    ws['H6'] = item.spec or None
    set_style(ws['H6'], 'detail_text')
    ws.merge_cells('A6:B6')
    ws.merge_cells('C6:E6')
    ws.merge_cells('F6:G6')
    ws.merge_cells('H6:J6')

    ws['A7'] = '日期范围'
    set_style(ws['A7'], 'meta_label')
    ws['C7'] = f'{start:%Y年%m月%d日} 至 {end:%Y年%m月%d日}'
    set_style(ws['C7'], 'detail_text')
    ws.merge_cells('A7:B7')
    ws.merge_cells('C7:J7')

    ws['A8'] = '交易明细'
    set_style(ws['A8'], 'title')
    ws.merge_cells('A8:J8')

    for col, header in enumerate(DETAIL_HEADERS, 1):
        cell = ws.cell(9, col)
        cell.value = header
        set_style(cell, 'detail_header')

    # opening row
    for col in range(1, DETAIL_COLS + 1):
        set_style(ws.cell(10, col), 'detail_text')
    ws.cell(10, 2).value = '上期结转'
    set_style(ws.cell(10, 2), 'detail_open_label')
    ws.cell(10, 4).value = item.opening_qty
    set_style(ws.cell(10, 4), 'detail_open_qty')
    ws.cell(10, 5).value = item.opening_amt
    set_style(ws.cell(10, 5), 'detail_open_amt')
    ws.cell(10, 8).value = item.opening_qty
    set_style(ws.cell(10, 8), 'detail_open_qty')
    ws.cell(10, 9).value = item.opening_amt
    set_style(ws.cell(10, 9), 'detail_open_amt')
    ws.cell(10, 10).value = '上期结转'
    set_style(ws.cell(10, 10), 'detail_text')

    row = 11
    for idx, tx in enumerate(item.current_rows):
        populate_detail_tx_row(ws, row, tx, idx)
        row += 1

    # total row
    for col in range(1, DETAIL_COLS + 1):
        set_style(ws.cell(row, col), 'detail_total_highlight')
    ws.cell(row, 1).value = None
    ws.cell(row, 2).value = '本期合计'
    ws.cell(row, 3).value = None
    ws.cell(row, 4).value = f'=+SUM(D11:D{row - 1})'
    ws.cell(row, 5).value = f'=+SUM(E11:E{row - 1})'
    ws.cell(row, 6).value = f'=+SUM(F11:F{row - 1})'
    ws.cell(row, 7).value = f'=+SUM(G11:G{row - 1})'
    ws.cell(row, 8).value = item.ending_qty
    ws.cell(row, 9).value = item.ending_amt
    ws.cell(row, 10).value = None

    set_style(ws.cell(row, 4), 'detail_total_green')
    set_style(ws.cell(row, 5), 'detail_total_orange')
    set_style(ws.cell(row, 6), 'detail_total_highlight')
    set_style(ws.cell(row, 7), 'detail_total_highlight')
    set_style(ws.cell(row, 8), 'detail_total_green')
    set_style(ws.cell(row, 9), 'detail_total_orange')
    set_style(ws.cell(row, 10), 'detail_total_purple')


# --- Validation -------------------------------------------------------------
def validate_generated_workbook(content: bytes) -> None:
    wb = load_workbook(io.BytesIO(content), data_only=False)
    if not wb.sheetnames or wb.sheetnames[0] != SUMMARY_SHEET_NAME:
        raise RuntimeError('Workbook validation failed: first sheet is not 汇总')
    if '_compat_check' in wb.sheetnames:
        raise RuntimeError('Workbook validation failed: unexpected temporary sheet _compat_check found')
    if DETAIL_TEMPLATE_SHEET in wb.sheetnames:
        raise RuntimeError('Workbook validation failed: hidden template sheet {} was not removed'.format(repr(DETAIL_TEMPLATE_SHEET)))
    summary = wb[SUMMARY_SHEET_NAME]
    if summary['A4'].value != '本期有动态':
        raise RuntimeError('Workbook validation failed: A4={}'.format(repr(summary['A4'].value)))
    expected_headers = SUMMARY_HEADERS
    actual_headers = [summary.cell(5, col).value for col in range(1, SUMMARY_COLS + 1)]
    if actual_headers != expected_headers:
        raise RuntimeError('Workbook validation failed: summary headers mismatch')
    inactive_title_rows = [
        row for row in range(1, summary.max_row + 1)
        if summary.cell(row, 1).value == '有结余 / 无动态 物料汇总'
    ]
    if len(inactive_title_rows) != 1:
        raise RuntimeError('Workbook validation failed: expected exactly one inactive section title, got {}'.format(repr(inactive_title_rows)))
    inactive_title_row = inactive_title_rows[0]
    if inactive_title_row < 6:
        raise RuntimeError(
            'Workbook validation failed: inactive section title row {} is before row 6'.format(
                inactive_title_row
            )
        )
    inactive_headers = [summary.cell(inactive_title_row + 1, col).value for col in range(1, SUMMARY_COLS + 1)]
    if inactive_headers != expected_headers:
        raise RuntimeError('Workbook validation failed: inactive summary headers mismatch')
    for ws in wb.worksheets:
        if len(ws.tables) != 0:
            raise RuntimeError('Workbook validation failed: unexpected tables found in sheet {}'.format(repr(ws.title)))
    if summary.max_column < SUMMARY_COLS:
        raise RuntimeError('Workbook validation failed: summary sheet is too narrow')
    active_section_rows = list(range(6, inactive_title_row))
    active_data_rows = [
        row for row in active_section_rows
        if any(
            summary.cell(row, col).value not in (None, '')
            for col in range(1, SUMMARY_COLS + 1)
        )
    ]
    blank_active_rows = [
        row for row in active_section_rows
        if row not in active_data_rows
    ]
    if blank_active_rows:
        raise RuntimeError(
            'Workbook validation failed: blank row found between active materials and inactive section: {}'.format(
                blank_active_rows
            )
        )
    expected_inactive_title_row = 6 + len(active_data_rows)
    if inactive_title_row != expected_inactive_title_row:
        raise RuntimeError(
            'Workbook validation failed: inactive section starts at row {}, expected row {}'.format(
                inactive_title_row,
                expected_inactive_title_row,
            )
        )
    if active_data_rows:
        first_active_row = active_data_rows[0]
        for col in (1, 2, 3):
            cell = summary.cell(first_active_row, col)
            if cell.value not in (None, '') and (cell.hyperlink is None or not cell.hyperlink.location):
                raise RuntimeError(
                    'Workbook validation failed: summary identity hyperlink missing at {}'.format(cell.coordinate)
                )
            if not cell.alignment.shrink_to_fit:
                raise RuntimeError(
                    'Workbook validation failed: shrink-to-fit missing at {}'.format(cell.coordinate)
                )
    if len(wb.sheetnames) > 1:
        detail = wb[wb.sheetnames[1]]
        if detail['F4'].value != '父物料组' or detail['F5'].value != '子物料组':
            raise RuntimeError('Workbook validation failed: detail group labels mismatch')
        if detail['A7'].value != '日期范围':
            raise RuntimeError('Workbook validation failed: detail date range row missing')
        return_cell = detail['I1']
        expected_return_location = sheet_location(SUMMARY_SHEET_NAME)
        formula_text = str(return_cell.value or '')
        if not formula_text.startswith('=HYPERLINK('):
            raise RuntimeError('Workbook validation failed: return-to-summary formula missing')
        if return_cell.hyperlink is None or return_cell.hyperlink.location != expected_return_location:
            raise RuntimeError('Workbook validation failed: return-to-summary hyperlink missing or invalid')


# --- ERPNext/Frappe entrypoint ----------------------------------------------------
def _read_connection_values() -> tuple[str, str, str]:
    """Read the ERPNext connection from the App-managed ERPNext 当前站点适配器.

    The URL is stored at f/erpnext/base_url and the credential at
    f/erpnext/api_token. The token value is never returned to the frontend or
    written to logs.
    """
    variable_errors: list[str] = []

    try:
        base_url = str(wmill.get_variable(BASE_URL_VARIABLE_PATH) or "").strip().rstrip("/")
    except Exception as exc:
        base_url = ""
        variable_errors.append(f"{BASE_URL_VARIABLE_PATH}: {exc}")

    try:
        token = str(wmill.get_variable(API_TOKEN_VARIABLE_PATH) or "").strip()
    except Exception as exc:
        token = ""
        variable_errors.append(f"{API_TOKEN_VARIABLE_PATH}: {exc}")

    missing: list[str] = []
    if not base_url:
        missing.append(BASE_URL_VARIABLE_PATH)
    if not token:
        missing.append(API_TOKEN_VARIABLE_PATH)
    if missing:
        detail = f"；读取错误：{' | '.join(variable_errors)}" if variable_errors else ""
        raise RuntimeError(
            "ERPNext连接配置不可用。请在ERPNext业务中心的“连接设置”页面完成配置："
            + "、".join(missing)
            + detail
        )

    return base_url, token, "frappe_local"

def main(
    period: str,
    company: str = DEFAULT_COMPANY,
    parent_group: str = DEFAULT_PARENT_GROUP,
) -> dict[str, Any]:
    """Generate and return a downloadable ERPNext material ledger workbook.

    The ERPNext connection is read dynamically from ERPNext 当前站点适配器. No AI is used.
    """
    start, end = parse_period(period)
    company = str(company or '').strip()
    parent_group = str(parent_group or '').strip()
    if not company:
        raise ValueError('company 不能为空')

    base_url, token, connection_source = _read_connection_values()
    erpnext = {"base_url": base_url, "api_token": token}

    print(f'[info] ERPNext connection source: {connection_source}')
    print(f'[info] Fetching ERPNext data from {start} to {end} ...')
    payload = fetch_erp_data(period, company, erpnext, start, end)

    print(f"[info] Building items (parent-group filter: '{parent_group or '无'}')...")
    items = build_items(payload, start, end, parent_group_filter=parent_group)
    active = [it for it in items if it.has_current]
    inactive = [it for it in items if not it.has_current]
    active.sort(
        key=lambda x: (x.parent_group or '', x.item_group or '', x.ending_amt, x.ending_qty, x.item_name, x.item_code),
        reverse=True,
    )
    inactive.sort(
        key=lambda x: (x.parent_group or '', x.item_group or '', x.ending_amt, x.ending_qty, x.item_name, x.item_code),
        reverse=True,
    )

    used_names: set[str] = set()
    for item in active:
        item.sheet_name = normalize_sheet_name(item.item_name, used_names)

    print(f'[info] Active: {len(active)}, Inactive: {len(inactive)}')

    wb = build_template_workbook()
    summary_ws = wb[SUMMARY_SHEET_NAME]
    build_summary_sheet(summary_ws, company, parent_group or '全部', start, end, active, inactive)
    for item in active:
        build_detail_sheet(wb, item, start, end)
    wb.active = 0
    if DETAIL_TEMPLATE_SHEET in wb.sheetnames:
        wb.remove(wb[DETAIL_TEMPLATE_SHEET])
    wb.calculation = CalcProperties(calcMode='auto', fullCalcOnLoad=True, forceFullCalc=True)

    short = infer_company_short(company)
    suffix = f'_{parent_group}' if parent_group else ''
    filename = f'{start:%Y%m%d}-{end:%Y%m%d}{short}{suffix}库存流水.xlsx'

    output = io.BytesIO()
    wb.save(output)
    content = output.getvalue()
    validate_generated_workbook(content)

    total_opening = round(sum(x.opening_amt for x in items), 2)
    total_inbound = round(sum(x.inbound_amt for x in items), 2)
    total_outbound = round(sum(x.outbound_amt for x in items), 2)
    total_ending = round(sum(x.ending_amt for x in items), 2)

    summary_rows = [
        {
            '统计周期': f'{start:%Y-%m-%d} 至 {end:%Y-%m-%d}',
            '公司': company,
            '父物料组': parent_group or '全部',
            '本期有动态物料数': len(active),
            '本期无动态物料数': len(inactive),
            '上期结转金额': total_opening,
            '本期入库金额': total_inbound,
            '本期出库金额': total_outbound,
            '本期结余金额': total_ending,
        }
    ]

    return {
        'ok': True,
        'backend_version': BACKEND_VERSION,
        'render_all': [
            {
                'markdown': (
                    '## ERPNext 库存流水已生成\n'
                    f'- 后端版本：{BACKEND_VERSION}\n'
                    f'- 日期范围：{start:%Y-%m-%d} 至 {end:%Y-%m-%d}\n'
                    f'- 公司：{company}\n'
                    f'- 父物料组：{parent_group or "全部"}\n'
                    f'- 本期有动态物料：{len(active)} 个\n'
                    f'- 本期结余金额：{total_ending:,.2f} 元'
                )
            },
            {'table-row-object': summary_rows},
            {
                'file': {
                    'content': base64.b64encode(content).decode('ascii'),
                    'filename': filename,
                }
            },
        ]
    }
