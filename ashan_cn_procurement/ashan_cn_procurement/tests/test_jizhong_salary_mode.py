import sys, os
from pathlib import Path
app_root = Path(__file__).resolve().parents[2]
inner_pkg = Path(__file__).resolve().parents[1]
if str(app_root) not in sys.path:
    sys.path.insert(0, str(app_root))
import ashan_cn_procurement
if str(inner_pkg) not in ashan_cn_procurement.__path__:
    ashan_cn_procurement.__path__.append(str(inner_pkg))

import unittest
from unittest.mock import patch
import sys
import types
from datetime import date, datetime

try:
    import frappe  # noqa: F401
except ModuleNotFoundError:
    frappe = types.ModuleType("frappe")
    frappe.throw = lambda message: (_ for _ in ()).throw(ValueError(message))
    frappe.db = types.SimpleNamespace()
    frappe.get_all = lambda *args, **kwargs: []
    frappe.get_list = lambda *args, **kwargs: []
    frappe.whitelist = lambda *args, **kwargs: (lambda fn: fn) if not (args and callable(args[0])) else args[0]
    frappe._ = lambda msg: msg
    sys.modules["frappe"] = frappe

    utils = types.ModuleType("frappe.utils")
    utils.cint = lambda value=0: int(float(value or 0))
    utils.flt = lambda value=0: float(value or 0)
    utils.today = lambda: date.today().isoformat()
    utils.now_datetime = lambda: datetime.now()
    utils.add_days = lambda dt, days: str(dt)
    utils.date_diff = lambda d1, d2: 0
    def _getdate(value=None):
        if value is None:
            return date.today()
        if isinstance(value, date):
            return value
        return datetime.strptime(str(value)[:10], "%Y-%m-%d").date()
    utils.getdate = _getdate
    sys.modules["frappe.utils"] = utils

import ashan_cn_procurement.services.jizhong_payroll_service as payroll_service

from ashan_cn_procurement.services.jizhong_payroll_service import (
    JIZHONG_SALARY_MODE_FIXED_NET,
    JIZHONG_SALARY_MODE_DYNAMIC_GROSS,
    JIZHONG_SALARY_MODE_FIXED_GROSS,
    JIZHONG_SALARY_MODE_DYNAMIC_NET,
    normalize_jizhong_salary_mode,
    is_jizhong_fixed_net_salary_mode,
    is_jizhong_post_tax_salary_mode,
    is_jizhong_dynamic_salary_mode,
    get_jizhong_salary_mode_label,
    derive_jizhong_gross_from_net,
    _jizhong_statutory_exclusion_reason,
)
from ashan_cn_procurement.services.jizhong_attendance_service import (
    _is_jizhong_attendance_required,
)


class TestJizhongSalaryMode(unittest.TestCase):
    """测试吉众税后固定工资语义、别名兼容、考勤免除与实发净额不变式。"""

    def test_canonical_normalization(self):
        """测试规范化映射：标准模式保留，旧版别名无损兼容。"""
        self.assertEqual(normalize_jizhong_salary_mode("税后固定工资"), JIZHONG_SALARY_MODE_FIXED_NET)
        self.assertEqual(normalize_jizhong_salary_mode("税后管理工资"), JIZHONG_SALARY_MODE_FIXED_NET)
        self.assertEqual(normalize_jizhong_salary_mode("税前动态工资"), JIZHONG_SALARY_MODE_DYNAMIC_GROSS)
        self.assertEqual(normalize_jizhong_salary_mode(""), JIZHONG_SALARY_MODE_DYNAMIC_GROSS)
        self.assertEqual(normalize_jizhong_salary_mode(None), JIZHONG_SALARY_MODE_DYNAMIC_GROSS)

    def test_fixed_net_predicate(self):
        """测试税后固定薪资谓词判断：标准与历史别名均返回 True。"""
        self.assertTrue(is_jizhong_fixed_net_salary_mode("税后固定工资"))
        self.assertTrue(is_jizhong_fixed_net_salary_mode("税后管理工资"))
        self.assertFalse(is_jizhong_fixed_net_salary_mode("税前动态工资"))
        self.assertFalse(is_jizhong_fixed_net_salary_mode("计件工资"))
        self.assertFalse(is_jizhong_fixed_net_salary_mode(None))

    def test_user_facing_label_helper(self):
        """测试前端展示文案助手：历史记录在 UI 上统一呈现为‘税后固定工资’。"""
        self.assertEqual(get_jizhong_salary_mode_label("税后固定工资"), "税后固定工资")
        self.assertEqual(get_jizhong_salary_mode_label("税后管理工资"), "税后固定工资")
        self.assertEqual(get_jizhong_salary_mode_label("税前动态工资"), "税前动态工资")
        self.assertEqual(get_jizhong_salary_mode_label(""), "税前动态工资")
        self.assertEqual(get_jizhong_salary_mode_label(None), "税前动态工资")

    def test_attendance_exemption_for_canonical_and_legacy(self):
        """测试考勤免除判定：标准税后固定工资与历史税后管理工资均享受免打卡豁免。"""
        canonical_profile = {
            "employee_no": "JZ-TEST-001",
            "employee_name": "测试人员甲",
            "employee_type": "其他",
            "salary_mode": "税后固定工资",
            "base_salary": 0.0,
            "fixed_salary": 8000.0,
        }
        legacy_profile = {
            "employee_no": "JZ-TEST-002",
            "employee_name": "测试人员乙",
            "employee_type": "其他",
            "salary_mode": "税后管理工资",
            "base_salary": 0.0,
            "fixed_salary": 8000.0,
        }
        standard_dynamic_profile = {
            "employee_no": "JZ-TEST-003",
            "employee_name": "测试人员丙",
            "employee_type": "正式工",
            "salary_mode": "税前动态工资",
            "base_salary": 3500.0,
            "fixed_salary": 0.0,
        }

        # 固定税后免考勤：_is_jizhong_attendance_required 应返回 False
        self.assertFalse(_is_jizhong_attendance_required(canonical_profile))
        self.assertFalse(_is_jizhong_attendance_required(legacy_profile))

        # 普通车间动态工时：必须打卡上传考勤
        self.assertTrue(_is_jizhong_attendance_required(standard_dynamic_profile))

    def test_final_net_equality_invariant(self):
        """测试闭式反推不变式：最终实发净薪分毫不差严格等于目标设定金额（精确到分）。"""
        test_amounts = [
            3000.00,
            5000.00,
            6000.00,
            7500.50,
            8888.88,
            10000.00,
            12345.67,
            15000.00,
            20000.00,
            35000.00,
        ]

        for target_net in test_amounts:
            for special_ded in [0.0, 500.0, 1024.50]:
                for add_ded in [0.0, 1000.0, 2000.0]:
                    res = derive_jizhong_gross_from_net(
                        target_net=target_net,
                        special_deductions_cur=special_ded,
                        tax_threshold=5000.0,
                        cur_additional_ded=add_ded,
                    )
                    gross = res["gross_salary"]
                    tax = res["tax_amount"]
                    final_net = round(gross - special_ded - tax, 2)

                    # 核心不变式校验：推导出的最终净薪必须 100% 精确等于设定的固定净薪
                    self.assertEqual(
                        final_net,
                        round(target_net, 2),
                        f"净薪不变式违背：target={target_net}, gross={gross}, tax={tax}, net={final_net}"
                    )
                    self.assertEqual(res["net_salary"], round(target_net, 2))

    def test_social_security_and_housing_fund_eligibility_are_independent(self):
        """One statutory contribution must not implicitly control the other."""

        class Policy(dict):
            __getattr__ = dict.get

        social_only = Policy(
            participate_social_security=1,
            participate_housing_fund=0,
            is_active=1,
            insurance_rule_description="",
        )
        with patch.object(
            payroll_service.frappe.db, "exists", return_value=True, create=True
        ), patch.object(
            payroll_service.frappe.db,
            "get_value",
            return_value=social_only,
            create=True,
        ):
            self.assertEqual(
                _jizhong_statutory_exclusion_reason("正式工", "social_security"),
                "",
            )
            self.assertEqual(
                _jizhong_statutory_exclusion_reason("正式工", "housing_fund"),
                "正式工不参加住房公积金",
            )

        housing_only = Policy(
            participate_social_security=0,
            participate_housing_fund=1,
            is_active=1,
            insurance_rule_description="",
        )
        with patch.object(
            payroll_service.frappe.db, "exists", return_value=True, create=True
        ), patch.object(
            payroll_service.frappe.db,
            "get_value",
            return_value=housing_only,
            create=True,
        ):
            self.assertEqual(
                _jizhong_statutory_exclusion_reason("正式工", "social_security"),
                "正式工不参加社会保险",
            )
            self.assertEqual(
                _jizhong_statutory_exclusion_reason("正式工", "housing_fund"),
                "",
            )


    def test_salary_adjustment_rules(self):
        """测试工资调整规则：正数增工资，负数扣工资；税后计入税后目标，税前计入税前应发。"""
        # 1. 税后固定工资：工资调整算入税后净额
        fixed_salary = 10500.0
        adj_pos = 500.0
        target_net_pos = round(fixed_salary + adj_pos, 2)
        self.assertEqual(target_net_pos, 11000.0)
        res_pos = derive_jizhong_gross_from_net(
            target_net=target_net_pos,
            special_deductions_cur=686.02,
            tax_threshold=5000.0,
            prev_gross=56239.4,
            prev_threshold=30000.0,
            prev_special_ded=4116.12,
            prev_additional_ded=0.0,
            prev_tax_paid=601.76,
        )
        verified_net = round(res_pos["gross_salary"] - 686.02 - res_pos["tax_amount"], 2)
        self.assertEqual(verified_net, 11000.0)

        # 负数工资调整：扣除工资
        adj_neg = -500.0
        target_net_neg = round(fixed_salary + adj_neg, 2)
        self.assertEqual(target_net_neg, 10000.0)
        res_neg = derive_jizhong_gross_from_net(
            target_net=target_net_neg,
            special_deductions_cur=686.02,
            tax_threshold=5000.0,
            prev_gross=56239.4,
            prev_threshold=30000.0,
            prev_special_ded=4116.12,
            prev_additional_ded=0.0,
            prev_tax_paid=601.76,
        )
        verified_net_neg = round(res_neg["gross_salary"] - 686.02 - res_neg["tax_amount"], 2)
        self.assertEqual(verified_net_neg, 10000.0)

        # 2. 税前工资：工资调整直接算入税前应发
        base_sal = 8000.0
        gross_pos = round(base_sal + adj_pos, 2)
        self.assertEqual(gross_pos, 8500.0)
        gross_neg = round(base_sal + adj_neg, 2)
        self.assertEqual(gross_neg, 7500.0)


    def test_jizhong_excel_export_sheets(self):
        """Verify Jizhong Cash Distribution (A4 + Envelope) and Payroll (36 cols) Excel exports."""
        import openpyxl
        import io
        import base64
        from ashan_cn_procurement.services.payroll_settlement_service import export_qifu_payroll_excel

        mock_items = [
            {
                "employee_no": "A0001",
                "employee_name": "陈亮",
                "base_salary": 10000.0,
                "house_rent_allowance": 0.0,
                "performance_salary": 0.0,
                "post_allowance": 500.0,
                "meal_unit_price": 15.0,
                "attendance_days": 25.0,
                "half_days": 0.0,
                "absent_days": 6.0,
                "basic_hours": 184.0,
                "overtime_regular_1_5": 31.0,
                "overtime_weekend_2_0": 16.0,
                "overtime_holiday_3_0": 0.0,
                "leave_compensatory_hours": 0.0,
                "meal_count": 0,
                "salary_basic_hours": 11247.88,
                "salary_overtime_1_5": 2703.49,
                "salary_overtime_2_0": 1860.47,
                "salary_overtime_3_0": 0.0,
                "salary_basic_subsidy": 0.0,
                "salary_performance": 0.0,
                "salary_post_allowance": 500.0,
                "salary_meal_subsidy": 0.0,
                "salary_adjustment": 0.0,
                "gross_salary": 11247.88,
                "hf_company_total": 126.0,
                "ss_company_total": 1411.66,
                "company_cost_total": 1537.66,
                "hf_person_total": 126.0,
                "ss_person_total": 560.02,
                "tax_amount": 61.86,
                "person_cost_total": 747.88,
                "net_salary": 10500.0,
                "cash_pay": 10500.0,
                "bills_100": 105,
                "bills_50": 0,
                "bills_10": 0,
                "bills_5": 0,
                "bills_1": 0,
            }
        ]

        with patch("ashan_cn_procurement.services.payroll_settlement_service.check_payroll_workbench_permission", return_value=True), \
             patch("ashan_cn_procurement.services.jizhong_payroll_service.get_jizhong_payroll_overview", return_value={"items": mock_items}), \
             patch.object(frappe, "get_list", create=True, return_value=[]), \
             patch("ashan_cn_procurement.services.ashan_holiday_service.get_month_workdays", return_value=23):

            # 1. Test Cash Distribution Export (A4 + Envelope)
            res_cash = export_qifu_payroll_excel(
                company="天津吉众科技有限公司",
                period_month="2026-07",
                sheet_type="jizhong_cash_distribution"
            )
            self.assertTrue(res_cash["success"])
            self.assertEqual(res_cash["filename"], "天津吉众科技有限公司_2026-07_工资发放表.xlsx")
            wb_cash = openpyxl.load_workbook(io.BytesIO(base64.b64decode(res_cash["file_base64"])))
            self.assertEqual(wb_cash.sheetnames, ["工资条-A4", "工资条-信封"])

            ws_a4 = wb_cash["工资条-A4"]
            self.assertEqual(ws_a4["A1"].value, "2026年07月工资确认、签收单")
            self.assertEqual(ws_a4.max_column, 24)
            self.assertEqual(ws_a4.cell(2, 1).value, "序号")
            self.assertEqual(ws_a4.cell(2, 24).value, "现金合计")
            self.assertEqual(ws_a4.cell(3, 2).value, "A0001")
            self.assertEqual(ws_a4.cell(3, 17).value, 10500.0)
            self.assertEqual(ws_a4.cell(4, 1).value, "合计")

            ws_env = wb_cash["工资条-信封"]
            self.assertEqual(ws_env.max_column, 8)
            self.assertEqual(ws_env.cell(1, 1).value, "序号")
            self.assertEqual(ws_env.cell(1, 8).value, "现金发放工资")
            self.assertEqual(ws_env.cell(2, 2).value, "A0001")
            self.assertEqual(ws_env.cell(2, 8).value, 10500.0)

            # 2. Test Payroll Export (36 cols)
            res_pr = export_qifu_payroll_excel(
                company="天津吉众科技有限公司",
                period_month="2026-07",
                sheet_type="jizhong_payroll"
            )
            self.assertTrue(res_pr["success"])
            self.assertEqual(res_pr["filename"], "天津吉众科技有限公司_2026-07_工资核定表.xlsx")
            wb_pr = openpyxl.load_workbook(io.BytesIO(base64.b64decode(res_pr["file_base64"])))
            self.assertEqual(wb_pr.sheetnames, ["本月工资核定表"])

            ws_pr = wb_pr["本月工资核定表"]
            self.assertEqual(ws_pr["A1"].value, "2026年07月员工工资核定表")
            self.assertIn("动态工作日=23天×8h=184h", ws_pr["A3"].value)
            self.assertEqual(ws_pr.max_column, 36)
            self.assertEqual(ws_pr.cell(4, 1).value, "序号")
            self.assertEqual(ws_pr.cell(4, 4).value, "薪资标准")
            self.assertEqual(ws_pr.cell(5, 4).value, "基本工资")
            self.assertEqual(ws_pr.cell(4, 36).value, "现金发放工资")
            self.assertEqual(ws_pr.cell(6, 2).value, "A0001")
            self.assertEqual(ws_pr.cell(6, 27).value, 11247.88)
            self.assertEqual(ws_pr.cell(7, 1).value, "合计")

if __name__ == "__main__":
    unittest.main()